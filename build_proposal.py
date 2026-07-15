"""Generate a ready-to-pitch business proposal for VideoClipper AI as a .docx file."""

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ---- Brand palette ---------------------------------------------------------
INK = RGBColor(0x11, 0x18, 0x27)       # near-black
ACCENT = RGBColor(0x6D, 0x28, 0xD9)    # violet
ACCENT_SOFT = RGBColor(0x8B, 0x5C, 0xF6)
MUTED = RGBColor(0x64, 0x74, 0x8B)     # slate
LIGHT = RGBColor(0xED, 0xE9, 0xFE)     # light violet fill
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

FONT = "Calibri"
HEAD_FONT = "Calibri"


def set_cell_bg(cell, hex_color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    m = OxmlElement("w:tcMar")
    for tag, val in (("top", top), ("bottom", bottom), ("start", left), ("end", right)):
        node = OxmlElement(f"w:{tag}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        m.append(node)
    tc_pr.append(m)


def remove_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "none")
        borders.append(el)
    tbl_pr.append(borders)


def shade_paragraph(paragraph, hex_color):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    p_pr.append(shd)


def add_bottom_border(paragraph, hex_color, size=8):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "6")
    bottom.set(qn("w:color"), hex_color)
    borders.append(bottom)
    p_pr.append(borders)


def style_run(run, size=11, bold=False, color=INK, italic=False, font=FONT):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color


doc = Document()

# Base style
normal = doc.styles["Normal"]
normal.font.name = FONT
normal.font.size = Pt(11)
normal.font.color.rgb = INK
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.15

for section in doc.sections:
    section.top_margin = Inches(0.9)
    section.bottom_margin = Inches(0.9)
    section.left_margin = Inches(0.95)
    section.right_margin = Inches(0.95)


def spacer(size=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    r = p.add_run("")
    r.font.size = Pt(size)
    return p


def heading(text, number=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(4)
    if number is not None:
        rn = p.add_run(f"{number}  ")
        style_run(rn, size=15, bold=True, color=ACCENT, font=HEAD_FONT)
    r = p.add_run(text)
    style_run(r, size=15, bold=True, color=INK, font=HEAD_FONT)
    add_bottom_border(p, "E5E7EB", size=6)
    return p


def subheading(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text)
    style_run(r, size=12, bold=True, color=ACCENT)
    return p


def body(text, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run(text)
    style_run(r, size=11, color=INK)
    return p


def bullet(text, bold_lead=None):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    if bold_lead:
        r = p.add_run(bold_lead)
        style_run(r, size=11, bold=True, color=INK)
        r2 = p.add_run(text)
        style_run(r2, size=11, color=INK)
    else:
        r = p.add_run(text)
        style_run(r, size=11, color=INK)
    return p


# ===========================================================================
# COVER
# ===========================================================================
spacer(40)

eyebrow = doc.add_paragraph()
eyebrow.alignment = WD_ALIGN_PARAGRAPH.LEFT
r = eyebrow.add_run("PROJECT PROPOSAL  ·  CONFIDENTIAL")
style_run(r, size=11, bold=True, color=ACCENT_SOFT)
eyebrow.paragraph_format.space_after = Pt(4)

title = doc.add_paragraph()
r = title.add_run("VideoClipper AI")
style_run(r, size=40, bold=True, color=INK)
title.paragraph_format.space_after = Pt(2)

sub = doc.add_paragraph()
r = sub.add_run("Turn one long video into a week of ready-to-post viral clips.")
style_run(r, size=16, color=MUTED)
sub.paragraph_format.space_after = Pt(10)
add_bottom_border(sub, "6D28D9", size=18)

spacer(14)

pitch = doc.add_paragraph()
r = pitch.add_run(
    "An AI clipping studio that ingests any long video or YouTube link and returns "
    "ranked, captioned, platform-ready 9:16 short clips \u2014 complete with titles, "
    "descriptions, hashtags and virality scores \u2014 in minutes, not hours."
)
style_run(r, size=12, color=INK)
pitch.paragraph_format.space_after = Pt(20)

# Prepared-for / by table
meta = doc.add_table(rows=4, cols=2)
meta.alignment = WD_TABLE_ALIGNMENT.LEFT
remove_table_borders(meta)
meta_rows = [
    ("Prepared for", "[Recipient name / company]"),
    ("Prepared by", "[Your name]"),
    ("Date", "[Date]"),
    ("Version", "1.0"),
]
for i, (k, v) in enumerate(meta_rows):
    c0, c1 = meta.rows[i].cells
    c0.width = Inches(1.6)
    c1.width = Inches(4.6)
    set_cell_margins(c0)
    set_cell_margins(c1)
    p0 = c0.paragraphs[0]
    rr = p0.add_run(k.upper())
    style_run(rr, size=9, bold=True, color=MUTED)
    p1 = c1.paragraphs[0]
    rr = p1.add_run(v)
    style_run(rr, size=11, bold=True, color=INK)

doc.add_page_break()

# ===========================================================================
# 1. EXECUTIVE SUMMARY
# ===========================================================================
heading("Executive Summary", 1)
body(
    "Short-form video is the single most effective way to grow an audience today, yet "
    "producing it is slow, manual and expensive. Creators, agencies and brands sit on "
    "hours of podcasts, webinars, interviews and streams that never get repurposed "
    "because editing them into clips takes real time and skill."
)
body(
    "VideoClipper AI removes that bottleneck. Upload a video or paste a YouTube link, and "
    "the platform automatically transcribes the audio, uses large language models to find "
    "the most engaging, self-contained moments, and renders polished vertical clips with "
    "animated captions, branding and background music \u2014 each one scored for viral "
    "potential and packaged with a ready-to-post title, description and hashtags."
)
body(
    "A working full-stack product already exists: a FastAPI backend orchestrating the AI "
    "and video pipeline, and a modern React studio interface. This proposal outlines the "
    "product, the opportunity, and a clear path to turn it into a revenue-generating "
    "service."
)

# Highlight stat strip
spacer(6)
stat = doc.add_table(rows=1, cols=3)
stat.alignment = WD_TABLE_ALIGNMENT.CENTER
remove_table_borders(stat)
stats = [
    ("Minutes", "From upload to downloadable clips"),
    ("1 \u2192 10+", "Clips generated per long video"),
    ("9:16", "Platform-ready for Shorts, Reels & TikTok"),
]
for i, (big, small) in enumerate(stats):
    cell = stat.rows[0].cells[i]
    set_cell_bg(cell, "EDE9FE")
    set_cell_margins(cell, top=140, bottom=140, left=120, right=120)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(big)
    style_run(r, size=20, bold=True, color=ACCENT)
    p2 = cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p2.add_run(small)
    style_run(r, size=9.5, color=MUTED)

# ===========================================================================
# 2. THE PROBLEM
# ===========================================================================
heading("The Problem", 2)
body("Every content team faces the same repurposing gap:")
bullet(" Editing a single clip \u2014 trimming, reframing to vertical, captioning and syncing \u2014 takes 30\u201360 minutes of skilled work.", "Editing is slow and manual.")
bullet(" Freelance editors and clipping agencies charge per clip, which quickly runs into hundreds of dollars per week for consistent output.", "It is expensive.")
bullet(" Deciding which 45 seconds of a two-hour stream will actually perform is guesswork for most creators.", "Choosing the right moment is hard.")
bullet(" Podcasts, webinars, courses and livestreams are packed with great moments that never get cut, so reach is left on the table.", "Great content goes to waste.")

# ===========================================================================
# 3. THE SOLUTION
# ===========================================================================
heading("The Solution", 3)
body(
    "VideoClipper AI is an end-to-end clipping studio that automates the entire workflow "
    "from raw footage to publish-ready shorts. The user stays in control \u2014 reviewing, "
    "re-timing and toggling edits \u2014 while the AI does the heavy lifting."
)

subheading("What makes it different")
bullet(" Every clip starts at the setup and ends at the payoff, so it makes sense on its own \u2014 not arbitrary fixed-length cuts.", "AI that understands complete thoughts. ")
bullet(" Each idea comes with virality, hook and content scores (0\u2013100) so users post the moments most likely to perform.", "Ranked by viral potential. ")
bullet(" Vertical 9:16 reframing, word-by-word animated captions, audio normalization, logo overlay and background music \u2014 all in one pass.", "Fully finished output. ")
bullet(" Titles, descriptions and hashtags are generated for every clip, ready to paste into TikTok, Reels or Shorts.", "Publish-ready packaging. ")

# ===========================================================================
# 4. HOW IT WORKS
# ===========================================================================
heading("How It Works", 4)
steps = [
    ("1", "Import", "Upload a local file (MP4, MOV, MKV, AVI) or paste a YouTube URL."),
    ("2", "Transcribe", "Audio is transcribed with word-level timestamps; long videos are chunked automatically."),
    ("3", "Select clips", "A large language model scans the transcript and picks the most shareable, self-contained moments."),
    ("4", "Review", "The user reviews ranked ideas with titles, scores, descriptions and hashtags, and fine-tunes in/out points."),
    ("5", "Export", "Clips are rendered as fast cuts or polished vertical edits with captions, branding and music."),
    ("6", "Download", "Grab individual MP4s or download every clip as a single ZIP."),
]
flow = doc.add_table(rows=len(steps), cols=2)
remove_table_borders(flow)
for i, (num, name, desc) in enumerate(steps):
    c0, c1 = flow.rows[i].cells
    c0.width = Inches(0.55)
    c1.width = Inches(5.9)
    set_cell_margins(c0, top=60, bottom=60)
    set_cell_margins(c1, top=60, bottom=60)
    set_cell_bg(c0, "6D28D9")
    p0 = c0.paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p0.add_run(num)
    style_run(r, size=14, bold=True, color=WHITE)
    p1 = c1.paragraphs[0]
    r = p1.add_run(name + " \u2014 ")
    style_run(r, size=11, bold=True, color=ACCENT)
    r = p1.add_run(desc)
    style_run(r, size=11, color=INK)

# ===========================================================================
# 5. KEY FEATURES
# ===========================================================================
heading("Key Features", 5)
features = [
    ("AI clip selection", "LLM-driven detection of hooks, payoffs, stories and debate moments, tuned for retention."),
    ("Virality scoring", "Every clip rated for virality, hook strength and content completeness."),
    ("Vertical reframing", "Automatic 9:16 conversion with a blurred-background fill for a native mobile look."),
    ("Animated captions", "Word-by-word highlighted subtitles, precisely synced to each clip."),
    ("Auto copywriting", "Ready-to-post titles, descriptions and hashtags generated per clip."),
    ("Branding controls", "Mask an existing corner logo and overlay your own in any corner."),
    ("Background music", "Upload a track and mix it softly under the original voice."),
    ("Audio normalization", "Even out loudness to a platform-friendly level automatically."),
    ("YouTube import", "Pull source video straight from a link \u2014 no manual download."),
    ("Batch export", "Download clips individually or all at once as a ZIP."),
]
ft = doc.add_table(rows=(len(features) + 1) // 2, cols=2)
remove_table_borders(ft)
for idx, (name, desc) in enumerate(features):
    row = idx // 2
    col = idx % 2
    cell = ft.rows[row].cells[col]
    cell.width = Inches(3.2)
    set_cell_margins(cell, top=60, bottom=60, left=100, right=140)
    p = cell.paragraphs[0]
    r = p.add_run("\u25B8 " + name)
    style_run(r, size=11, bold=True, color=ACCENT)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(4)
    r = p2.add_run(desc)
    style_run(r, size=10, color=MUTED)

# ===========================================================================
# 6. TECHNOLOGY
# ===========================================================================
heading("Technology", 6)
body(
    "VideoClipper AI is a real, working full-stack application \u2014 not a concept. It is "
    "built on a modern, scalable stack that keeps AI costs low and rendering fast."
)
tech = doc.add_table(rows=6, cols=2)
tech.alignment = WD_TABLE_ALIGNMENT.LEFT
remove_table_borders(tech)
tech_rows = [
    ("Frontend", "React + TypeScript + Vite \u2014 a responsive three-panel studio UI with live progress."),
    ("Backend", "Python FastAPI service exposing a clean REST API for the full pipeline."),
    ("Transcription", "Groq Whisper with word-level timestamps and automatic chunking for long media."),
    ("AI selection", "Groq Llama 3.3 70B for clip selection, with an optional Gemini Flash fallback."),
    ("Video engine", "FFmpeg for frame-accurate cutting, reframing, captions, mixing and encoding."),
    ("Media sources", "yt-dlp for YouTube ingestion; automatic cleanup of temporary files."),
]
for i, (k, v) in enumerate(tech_rows):
    c0, c1 = tech.rows[i].cells
    c0.width = Inches(1.5)
    c1.width = Inches(4.9)
    set_cell_margins(c0, top=60, bottom=60)
    set_cell_margins(c1, top=60, bottom=60)
    if i % 2 == 0:
        set_cell_bg(c0, "F5F3FF")
        set_cell_bg(c1, "F5F3FF")
    p0 = c0.paragraphs[0]
    r = p0.add_run(k)
    style_run(r, size=11, bold=True, color=INK)
    p1 = c1.paragraphs[0]
    r = p1.add_run(v)
    style_run(r, size=10.5, color=INK)

# ===========================================================================
# 7. MARKET OPPORTUNITY
# ===========================================================================
heading("Market Opportunity", 7)
body(
    "Short-form video dominates attention across TikTok, YouTube Shorts and Instagram "
    "Reels, and every platform rewards consistent posting. The demand for tools that make "
    "clipping effortless is large and growing \u2014 proven by the rapid rise of AI clipping "
    "products in the same category."
)
subheading("Who it is for")
bullet(" Podcasters, streamers and YouTubers who need daily shorts from long episodes.", "Creators. ")
bullet(" Social and video agencies producing clips for multiple clients at scale.", "Agencies. ")
bullet(" Marketing teams repurposing webinars, talks and product videos.", "Brands & B2B. ")
bullet(" Course creators and speakers turning long sessions into promo clips.", "Educators. ")

# ===========================================================================
# 8. BUSINESS MODEL
# ===========================================================================
heading("Business Model", 8)
body(
    "The natural model is usage-based SaaS: subscription tiers priced on processing "
    "minutes and clip volume, with higher tiers unlocking branding, music and priority "
    "rendering. Illustrative tiers:"
)
price = doc.add_table(rows=4, cols=3)
price.alignment = WD_TABLE_ALIGNMENT.CENTER
remove_table_borders(price)
plans = [
    ("Starter", "For solo creators", ["Monthly upload quota", "Vertical clips + captions", "Auto titles & hashtags"]),
    ("Pro", "For active creators", ["Higher quota", "Branding, logo & music", "Priority rendering"]),
    ("Agency", "For teams & clients", ["Large / pooled quota", "Multiple brand kits", "API & bulk export"]),
]
# header row
hdr = price.rows[0].cells
for i, (name, tag, _feat) in enumerate(plans):
    cell = hdr[i]
    set_cell_bg(cell, "6D28D9")
    set_cell_margins(cell, top=100, bottom=100)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(name)
    style_run(r, size=13, bold=True, color=WHITE)
    p2 = cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p2.add_run(tag)
    style_run(r, size=9, color=RGBColor(0xE9, 0xD5, 0xFF))
# feature rows
for r_i in range(1, 4):
    for c_i, (_n, _t, feats) in enumerate(plans):
        cell = price.rows[r_i].cells[c_i]
        set_cell_margins(cell, top=50, bottom=50)
        if r_i % 2 == 1:
            set_cell_bg(cell, "F5F3FF")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(feats[r_i - 1])
        style_run(r, size=10, color=INK)
spacer(4)
body(
    "Additional revenue paths include pay-as-you-go credits, white-label licensing for "
    "agencies, and API access for platforms that want clipping built in.",
    space_after=2,
)

# ===========================================================================
# 9. COMPETITIVE ADVANTAGE
# ===========================================================================
heading("Competitive Advantage", 9)
bullet(" A working full-stack product already exists and can be demoed today.", "Already built. ")
bullet(" Groq inference keeps transcription and selection fast and inexpensive, protecting margins.", "Cost-efficient AI. ")
bullet(" Clips are cut on complete thoughts, not blind fixed lengths \u2014 higher quality output.", "Smarter clipping. ")
bullet(" Reframing, captions, branding, music and copywriting ship in a single pipeline.", "All-in-one. ")
bullet(" Provider fallback (Groq \u2192 Gemini) keeps the service resilient to outages and rate limits.", "Resilient. ")

# ===========================================================================
# 10. ROADMAP
# ===========================================================================
heading("Roadmap", 10)
roadmap = doc.add_table(rows=4, cols=2)
remove_table_borders(roadmap)
phases = [
    ("Phase 1 \u2014 Live", "Working studio: import, transcribe, AI selection, vertical export, captions, branding, music, ZIP."),
    ("Phase 2 \u2014 Next", "User accounts, cloud storage, billing & subscription tiers, project history."),
    ("Phase 3 \u2014 Scale", "Background render queue, team workspaces, brand kits, scheduling & direct publishing."),
    ("Phase 4 \u2014 Expand", "Public API, mobile companion, multi-language captions, A/B thumbnail & title testing."),
]
for i, (name, desc) in enumerate(phases):
    c0, c1 = roadmap.rows[i].cells
    c0.width = Inches(1.7)
    c1.width = Inches(4.7)
    set_cell_margins(c0, top=70, bottom=70)
    set_cell_margins(c1, top=70, bottom=70)
    set_cell_bg(c0, "EDE9FE")
    p0 = c0.paragraphs[0]
    r = p0.add_run(name)
    style_run(r, size=10.5, bold=True, color=ACCENT)
    p1 = c1.paragraphs[0]
    r = p1.add_run(desc)
    style_run(r, size=10.5, color=INK)

# ===========================================================================
# 11. THE ASK / NEXT STEPS
# ===========================================================================
heading("The Ask & Next Steps", 11)
body(
    "I am looking for [an investment / a partnership / your first pilot customers] to take "
    "VideoClipper AI from a working product to a launched service. The immediate priorities "
    "are user accounts, billing and cloud rendering so the platform can onboard paying "
    "customers."
)
subheading("Proposed next steps")
bullet(" A 15-minute live walkthrough of the working product on a real video.", "Live demo. ")
bullet(" Run your own footage through the pipeline and review the output.", "Pilot. ")
bullet(" Align on scope, timeline and terms for the next phase.", "Agreement. ")

spacer(10)
cta = doc.add_paragraph()
shade_paragraph(cta, "111827")
cta.paragraph_format.space_before = Pt(6)
cta.paragraph_format.space_after = Pt(6)
cta.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = cta.add_run("Let\u2019s turn your long videos into a clip engine. \u2014 Book a demo today.")
style_run(r, size=13, bold=True, color=WHITE)

# ===========================================================================
# CONTACT FOOTER
# ===========================================================================
spacer(14)
contact = doc.add_paragraph()
add_bottom_border(contact, "E5E7EB", size=6)
contact.paragraph_format.space_after = Pt(6)

c = doc.add_paragraph()
r = c.add_run("Contact\n")
style_run(r, size=11, bold=True, color=ACCENT)
r = c.add_run("[Your name]  ·  [Email]  ·  [Phone]  ·  [Website / demo link]")
style_run(r, size=10.5, color=MUTED)

out = r"D:\AIProjects Desktop\clipper\VideoClipper_AI_Proposal.docx"
doc.save(out)
print("Saved:", out)
